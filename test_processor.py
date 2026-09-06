import os
import openpyxl
from processor import (
    parse_input_workbook,
    generate_sf9_front_workbook,
    generate_sf9_back_workbook,
    generate_sf9_zip_bundle,
    wrap_remarks_text
)

# 1. Test text wrapping utility
long_remark = "The student shows outstanding dedication to all academic subjects and demonstrates consistent leadership in classroom discussions and collaborative team activities."
wrapped = wrap_remarks_text(long_remark, width=56, max_lines=3)
print("Wrapped lines count:", len(wrapped))
for i, line in enumerate(wrapped, 1):
    print(f"Line {i} ({len(line)} chars): {line}")
    assert len(line) <= 56
assert len(wrapped) >= 2

# 2. Create a test workbook with 3 students (Testing 2-Up pairing with odd count: Sheet 1 = 1-2, Sheet 2 = 3)
wb_in = openpyxl.Workbook()
ws_in = wb_in.active

headers = [
    'School_Year', 'Name', 'Age', 'Sex', 'LRN', 'Grade', 'Section',
    'Filipino_T1', 'Filipino_T2', 'Filipino_T3',
    'English_T1', 'English_T2', 'English_T3',
    'Math_T1', 'Math_T2', 'Math_T3',
    'Science_T1', 'Science_T2', 'Science_T3',
    'AP_T1', 'AP_T2', 'AP_T3',
    'Values_T1', 'Values_T2', 'Values_T3',
    'TLE_T1', 'TLE_T2', 'TLE_T3',
    'MAPEH_T1', 'MAPEH_T2', 'MAPEH_T3',
    'MusicArts_T1', 'MusicArts_T2', 'MusicArts_T3',
    'PEHealth_T1', 'PEHealth_T2', 'PEHealth_T3',
    'Research1_T1', 'Research1_T2', 'Research1_T3',
    'MathInv_T1', 'MathInv_T2', 'MathInv_T3',
    'Days_Jun', 'Days_Jul', 'Days_Aug', 'Days_Sep', 'Days_Oct', 'Days_Nov', 'Days_Dec', 'Days_Jan', 'Days_Feb', 'Days_Mar', 'Days_Apr',
    'Present_Jun', 'Present_Jul', 'Present_Aug', 'Present_Sep', 'Present_Oct', 'Present_Nov', 'Present_Dec', 'Present_Jan', 'Present_Feb', 'Present_Mar', 'Present_Apr',
    'Absent_Jun', 'Absent_Jul', 'Absent_Aug', 'Absent_Sep', 'Absent_Oct', 'Absent_Nov', 'Absent_Dec', 'Absent_Jan', 'Absent_Feb', 'Absent_Mar', 'Absent_Apr',
    'Remarks_T1', 'Remarks_T2', 'Remarks_T3',
    'Admitted_To_Grade', 'Eligible_For', 'Admitted_In', 'Transfer_Date'
]
ws_in.append(headers)

# Student 1: Navarro
ws_in.append([
    '2026-2027', 'Navarro, Juan', 14, 'Male', '109876543210', '8', 'Rizal',
    88, 89, 90, 91, 92, 93, 85, 87, 89, 86, 88, 90, 90, 91, 92, 95, 96, 97, 88, 89, 90, 89, 90, 91, 89, 90, 91, 89, 90, 91, 87, 88, 89, 86, 87, 88,
    20, 22, 21, 21, 22, 20, 15, 21, 20, 22, 20,
    20, 22, 20, 21, 22, 20, 15, 20, 20, 22, 20,
    0, 0, 1, 0, 0, 0, 0, 1, 0, 0, 0,
    long_remark, 'Maintained high academic standing', 'Promoted to Grade 9',
    'Grade 9', 'Grade 9', 'PGCHS', '2027-04-15'
])

# Student 2: Cedro
ws_in.append([
    '2026-2027', 'Cedro, Maria', 14, 'Female', '109876543211', '8', 'Rizal',
    92, 94, 95, 90, 92, 94, 94, 95, 96, 93, 95, 96, 92, 93, 95, 96, 97, 98, 90, 92, 93, 92, 93, 94, 91, 92, 93, 93, 94, 95, 95, 96, 97, 94, 95, 96,
    20, 22, 21, 21, 22, 20, 15, 21, 20, 22, 20,
    20, 22, 21, 21, 22, 20, 15, 21, 20, 22, 20,
    0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0,
    'Active participant in class activities', 'Consistent top rank', 'Promoted with Honors',
    'Grade 9', 'Grade 9', 'PGCHS', '2027-04-15'
])

# Student 3: Perez (Odd count)
ws_in.append([
    '2026-2027', 'Perez, Carlos', 13, 'Male', '109876543212', '8', 'Rizal',
    85, 86, 87, 88, 89, 90, 85, 86, 87, 87, 88, 89, 88, 89, 90, 90, 91, 92, 86, 87, 88, 87, 88, 89, 87, 88, 89, 87, 88, 89, 85, 86, 87, 85, 86, 87,
    20, 22, 21, 21, 22, 20, 15, 21, 20, 22, 20,
    19, 22, 21, 20, 22, 20, 15, 21, 20, 22, 20,
    1, 0, 0, 1, 0, 0, 0, 0, 0, 0, 0,
    'Satisfactory performance', 'Very cooperative', 'Promoted',
    'Grade 9', 'Grade 9', 'PGCHS', '2027-04-15'
])

test_path = "test_2up_input.xlsx"
wb_in.save(test_path)

students = parse_input_workbook(test_path)
assert len(students) == 3

# 3. Test FRONT Workbook
wb_front, front_fname = generate_sf9_front_workbook(students)
print("FRONT Sheets:", wb_front.sheetnames)
assert len(wb_front.sheetnames) == 2
assert "1-2 FRONT Navarro - Cedro" == wb_front.sheetnames[0]
assert "3 FRONT Perez" == wb_front.sheetnames[1]

# Check Sheet 1 FRONT: Left Card (Navarro) & Right Card (Cedro)
ws_f1 = wb_front[wb_front.sheetnames[0]]

# Image check: All 4 images preserved
assert len(ws_f1._images) == 4
print("[PASS] FRONT Sheet 1 has 4 images preserved!")

# Left Card checks
assert ws_f1['C14'].value == 'Navarro, Juan'
assert ws_f1['F14'].value == 14
assert ws_f1['D24'].value == 88
assert ws_f1['D24'].alignment.horizontal == 'center'
assert ws_f1['D24'].alignment.vertical == 'center'
assert ws_f1['D31'].value == 89 # MAPEH
assert ws_f1['G37'].value is not None # General Average Final

# Right Card checks
assert ws_f1['L14'].value == 'Cedro, Maria'
assert ws_f1['O14'].value == 14
assert ws_f1['M24'].value == 92
assert ws_f1['M24'].alignment.horizontal == 'center'
assert ws_f1['M24'].alignment.vertical == 'center'
assert ws_f1['M31'].value == 92 # MAPEH
assert ws_f1['P37'].value is not None # General Average Final
print("[PASS] FRONT 2-Up Sheet 1 verified!")

# 4. Test BACK Workbook
wb_back, back_fname = generate_sf9_back_workbook(students)
print("BACK Sheets:", wb_back.sheetnames)
assert len(wb_back.sheetnames) == 2
assert "1-2 BACK Navarro - Cedro" == wb_back.sheetnames[0]
assert "3 BACK Perez" == wb_back.sheetnames[1]

ws_b1 = wb_back[wb_back.sheetnames[0]]

# Left Card checks (Student 2: Cedro on Left for back-to-back printing)
assert ws_b1['B1'].value == 'Cedro, Maria'
assert ws_b1['C4'].value == 20
assert ws_b1['C4'].alignment.horizontal == 'center'
assert ws_b1['N4'].value == "=SUM(C4:M4)"
assert ws_b1['C9'].value == 'Active participant in class activities'
assert ws_b1['E31'].value == 'Grade 9'

# Right Card checks (Student 1: Navarro on Right for back-to-back printing)
assert ws_b1['Q1'].value == 'Navarro, Juan'
assert ws_b1['R4'].value == 20
assert ws_b1['R4'].alignment.horizontal == 'center'
assert ws_b1['AC4'].value == "=SUM(R4:AB4)"
# Check remarks wrapping in R9, R10
assert ws_b1['R9'].value is not None
assert ws_b1['R10'].value is not None
print("Navarro Term 1 Remark R9:", repr(ws_b1['R9'].value))
print("Navarro Term 1 Remark R10:", repr(ws_b1['R10'].value))
assert len(ws_b1['R9'].value) <= 56
assert ws_b1['T31'].value == 'Grade 9'
print("[PASS] BACK 2-Up Sheet 1 verified (Duplex alignment: Student 2 Left, Student 1 Right)!")

# Odd count page check (Perez alone on Sheet 2 -> Right side on BACK page)
ws_b2 = wb_back[wb_back.sheetnames[1]]
assert ws_b2['B1'].value is None
assert ws_b2['Q1'].value == 'Perez, Carlos'
print("[PASS] BACK 2-Up Sheet 2 verified (odd count page B1 empty, Q1 set for duplex)!")

# 5. Test user-provided computed final grades and general average
wb_computed = openpyxl.Workbook()
ws_comp = wb_computed.active
ws_comp.append([
    'School_Year', 'Name', 'Age', 'Sex', 'LRN', 'Grade', 'Section',
    'Filipino_T1', 'Filipino_T2', 'Filipino_T3', 'Filipino_Final', 'Filipino_Remarks',
    'General_Average'
])
ws_comp.append([
    '2026-2027', 'Santos, Anna', 14, 'Female', '123456789012', '7', 'Diamond',
    90, 92, 94, 93, 'Passed with Merit',
    95
])
comp_path = "test_computed_input.xlsx"
wb_computed.save(comp_path)

students_comp = parse_input_workbook(comp_path)
assert len(students_comp) == 1
assert students_comp[0].get('filipino_final') == 93
assert students_comp[0].get('filipino_remarks') == 'Passed with Merit'
assert students_comp[0].get('genavg_final') == 95

# 6. Test partial term entry (Final grade only reflected when all 3 terms are present)
wb_partial = openpyxl.Workbook()
ws_part = wb_partial.active
ws_part.append([
    'School_Year', 'Name', 'Age', 'Sex', 'LRN', 'Grade', 'Section',
    'Filipino_T1', 'Filipino_T2', 'English_T1'
])
ws_part.append([
    '2026-2027', 'Garcia, Miguel', 13, 'Male', '987654321098', '7', 'Ruby',
    89.6, 91.4, 88.2
])
part_path = "test_partial_input.xlsx"
wb_partial.save(part_path)

students_part = parse_input_workbook(part_path)
assert len(students_part) == 1

wb_part_front, _ = generate_sf9_front_workbook(students_part)
ws_pf = wb_part_front.active
# Check whole-number rounding on the generated card
assert ws_pf['D24'].value == 90  # 89.6 rounded to whole number 90
assert ws_pf['E24'].value == 91  # 91.4 rounded to whole number 91
assert ws_pf['F24'].value is None
assert ws_pf['D25'].value == 88  # 88.2 rounded to whole number 88
# Final Grade formula requires COUNT=3
assert ws_pf['G24'].value == '=IF(COUNT(D24:F24)=3,ROUND(AVERAGE(D24:F24),0),"")'
# 7. Test MAPEH independence and empty fields preservation
wb_empty = openpyxl.Workbook()
ws_emp = wb_empty.active
ws_emp.append(['School_Year', 'Name', 'LRN', 'MusicArts_T1', 'PEHealth_T1'])
ws_emp.append(['2026-2027', 'Luna, Antonio', '112233445566', 90, 92])
emp_path = "test_empty_input.xlsx"
wb_empty.save(emp_path)

students_emp = parse_input_workbook(emp_path)
wb_emp_front, _ = generate_sf9_front_workbook(students_emp)
ws_ef = wb_emp_front.active
# MAPEH T1 (D31) should NOT be auto-filled from Music (D32) and PE (D33); it remains None/empty
assert ws_ef['D31'].value is None
assert ws_ef['D32'].value == 90  # Music T1
assert ws_ef['D33'].value == 92  # PE T1
# Age, Sex, Section were omitted, so they remain None/empty
assert ws_ef['F14'].value is None
assert ws_ef['H14'].value is None
assert ws_ef['H15'].value is None
# General average core subjects range explicitly excludes rows 32 and 33
assert 'D24:D31,D34:D35' in ws_ef['D37'].value
print("[PASS] MAPEH independence and empty fields preservation verified!")

print("\nALL 2-UP PROCESSOR TESTS PASSED!")
