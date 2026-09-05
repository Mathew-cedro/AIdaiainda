import os
import io
import zipfile
import openpyxl
from fastapi.testclient import TestClient
from main import app

client = TestClient(app)

def test_health():
    response = client.get("/api/health")
    assert response.status_code == 200
    res = response.json()
    assert res["status"] == "ok"
    assert res["front_template_exists"] is True
    assert res["back_template_exists"] is True
    print("[PASS] /api/health passed")

def test_template_download():
    response = client.get("/api/template")
    assert response.status_code == 200
    assert "spreadsheetml" in response.headers["content-type"]
    print("[PASS] /api/template passed")

def test_preview_and_conversions():
    # Create a dummy workbook with 2 students
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append([
        'School_Year', 'Name', 'Age', 'Sex', 'LRN', 'Grade', 'Section',
        'Filipino_T1', 'Filipino_T2', 'Filipino_T3',
        'English_T1', 'English_T2', 'English_T3',
        'Math_T1', 'Math_T2', 'Math_T3',
        'Science_T1', 'Science_T2', 'Science_T3',
        'AP_T1', 'AP_T2', 'AP_T3',
        'Values_T1', 'Values_T2', 'Values_T3',
        'TLE_T1', 'TLE_T2', 'TLE_T3',
        'MAPEH_T1', 'MAPEH_T2', 'MAPEH_T3',
        'MusicArts_T1', 'MusicArts_T2', 'MusicArts_T3',
        'PEHealth_T1', 'PEHealth_T2', 'PEHealth_T3',
        'Research1_T1', 'Research1_T2', 'Research1_T3',
        'MathInv_T1', 'MathInv_T2', 'MathInv_T3',
        'Days_Jun', 'Days_Jul', 'Days_Aug', 'Days_Sep', 'Days_Oct', 'Days_Nov', 'Days_Dec', 'Days_Jan', 'Days_Feb', 'Days_Mar', 'Days_Apr',
        'Present_Jun', 'Present_Jul', 'Present_Aug', 'Present_Sep', 'Present_Oct', 'Present_Nov', 'Present_Dec', 'Present_Jan', 'Present_Feb', 'Present_Mar', 'Present_Apr',
        'Absent_Jun', 'Absent_Jul', 'Absent_Aug', 'Absent_Sep', 'Absent_Oct', 'Absent_Nov', 'Absent_Dec', 'Absent_Jan', 'Absent_Feb', 'Absent_Mar', 'Absent_Apr',
        'Remarks_T1', 'Remarks_T2', 'Remarks_T3',
        'Admitted_To_Grade', 'Eligible_For', 'Admitted_In', 'Transfer_Date'
    ])

    ws.append([
        '2026-2027', 'Navarro, Juan', 14, 'Male', '111111111111', '8', 'St. Matthew',
        85, 86, 87, 88, 89, 90, 91, 92, 93, 89, 90, 91, 88, 89, 90, 92, 93, 94, 86, 87, 88, 88, 89, 90, 88, 89, 90, 88, 89, 90, 87, 88, 89, 86, 87, 88,
        20, 21, 22, 20, 21, 22, 15, 20, 21, 22, 20,
        20, 21, 22, 20, 21, 22, 15, 20, 21, 22, 20,
        0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0,
        'Excellent work in class activities', 'Active in class', 'Promoted',
        'Grade 9', 'Grade 9', 'PGCHS', '2027-04-15'
    ])

    ws.append([
        '2026-2027', 'Cedro, Maria', 13, 'Female', '222222222222', '8', 'St. Matthew',
        90, 91, 92, 92, 93, 94, 95, 96, 97, 94, 95, 96, 93, 94, 95, 97, 98, 99, 91, 92, 93, 94, 95, 96, 94, 95, 96, 94, 95, 96, 96, 97, 98, 95, 96, 97,
        20, 21, 22, 20, 21, 22, 15, 20, 21, 22, 20,
        20, 21, 21, 20, 21, 22, 15, 20, 21, 22, 20,
        0, 0, 1, 0, 0, 0, 0, 0, 0, 0, 0,
        'Top Performer in all subjects', 'High honors', 'Promoted with Honors',
        'Grade 9', 'Grade 9', 'PGCHS', '2027-04-15'
    ])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    file_bytes = buf.getvalue()

    # 1. Preview
    res_preview = client.post(
        "/api/preview",
        files={"file": ("test_batch.xlsx", file_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    )
    assert res_preview.status_code == 200
    data = res_preview.json()
    assert data["success"] is True
    assert data["count"] == 2
    assert data["front_filename"] == "SF9_FRONT_Navarro-Cedro.xlsx"
    assert data["back_filename"] == "SF9_BACK_Navarro-Cedro.xlsx"
    assert data["zip_filename"] == "SF9_Complete_Navarro-Cedro.zip"
    print("[PASS] /api/preview passed")

    # 2. Convert FRONT (2-Up -> 1 sheet with 2 students)
    res_front = client.post(
        "/api/convert/front",
        files={"file": ("test_batch.xlsx", file_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data={"overrides": '{"adviser_name": "Mrs. Reyes", "school_head": "Dr. Principal"}'}
    )
    assert res_front.status_code == 200
    assert "SF9_FRONT_Navarro-Cedro.xlsx" in res_front.headers["content-disposition"]
    wb_f = openpyxl.load_workbook(io.BytesIO(res_front.content))
    assert len(wb_f.sheetnames) == 1
    assert "1-2 FRONT Navarro - Cedro" == wb_f.sheetnames[0]
    ws_f = wb_f["1-2 FRONT Navarro - Cedro"]
    # Left Card
    assert ws_f["C14"].value == "Navarro, Juan"
    assert ws_f["D24"].value == 85
    assert ws_f["D24"].alignment.horizontal == 'center'
    # Right Card
    assert ws_f["L14"].value == "Cedro, Maria"
    assert ws_f["M24"].value == 90
    assert ws_f["M24"].alignment.horizontal == 'center'
    print("[PASS] /api/convert/front passed with 2-Up sheet")

    # 3. Convert BACK (2-Up -> 1 sheet with 2 students)
    res_back = client.post(
        "/api/convert/back",
        files={"file": ("test_batch.xlsx", file_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data={"overrides": '{"adviser_name": "Mrs. Reyes", "school_head": "Dr. Principal"}'}
    )
    assert res_back.status_code == 200
    assert "SF9_BACK_Navarro-Cedro.xlsx" in res_back.headers["content-disposition"]
    wb_b = openpyxl.load_workbook(io.BytesIO(res_back.content))
    assert len(wb_b.sheetnames) == 1
    assert "1-2 BACK Navarro - Cedro" == wb_b.sheetnames[0]
    ws_b = wb_b["1-2 BACK Navarro - Cedro"]
    # Left Card
    assert ws_b["B1"].value == "Navarro, Juan"
    assert ws_b["C4"].value == 20
    assert ws_b["C4"].alignment.horizontal == 'center'
    assert ws_b["N4"].value == "=SUM(C4:M4)"
    assert ws_b["C9"].value == "Excellent work in class activities"
    assert ws_b["J34"].value == "Mrs. Reyes"
    # Right Card
    assert ws_b["Q1"].value == "Cedro, Maria"
    assert ws_b["R4"].value == 20
    assert ws_b["R4"].alignment.horizontal == 'center'
    assert ws_b["AC4"].value == "=SUM(R4:AB4)"
    assert ws_b["R9"].value == "Top Performer in all subjects"
    assert ws_b["Y34"].value == "Mrs. Reyes"
    print("[PASS] /api/convert/back passed with 2-Up sheet, B1/Q1 student names, and overrides")

    # 4. Convert ZIP
    res_zip = client.post(
        "/api/convert/zip",
        files={"file": ("test_batch.xlsx", file_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    )
    assert res_zip.status_code == 200
    assert "SF9_Complete_Navarro-Cedro.zip" in res_zip.headers["content-disposition"]
    zf = zipfile.ZipFile(io.BytesIO(res_zip.content))
    namelist = zf.namelist()
    assert "SF9_FRONT_Navarro-Cedro.xlsx" in namelist
    assert "SF9_BACK_Navarro-Cedro.xlsx" in namelist
    print("[PASS] /api/convert/zip passed with both files included")

if __name__ == "__main__":
    test_health()
    test_template_download()
    test_preview_and_conversions()
    print("\nALL ENDPOINT TESTS PASSED SUCCESSFULLY!")
