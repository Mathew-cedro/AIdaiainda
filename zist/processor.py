import os
import re
import io
import textwrap
import zipfile
import openpyxl
from openpyxl.styles import Alignment, Font
from typing import List, Dict, Any, Tuple, Optional

FRONT_TEMPLATE_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "SF9 FINAL FORMAT-FRONT.xlsx")
BACK_TEMPLATE_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "SF9 FINAL FORMAT-BACK.xlsx")

# Alignment and Font helpers
CENTER_ALIGN = Alignment(horizontal='center', vertical='center')
LEFT_CENTER_ALIGN = Alignment(horizontal='left', vertical='center')
NAME_HEADER_FONT = Font(name='Arial', size=11, bold=True, color='000000')

# Flexible header normalization map
HEADER_MAP = {
    # Profile
    'school_year': ['school_year', 'schoolyear', 'sy', 's.y.', 's.y', 'taong_panuruan'],
    'name': ['name', 'student_name', 'studentname', 'learners_name', 'learner_name', 'fullname', 'full_name', 'pangalan', 'student', 'learnersname', 'complete_name', 'student_full_name', 'pangalan_ng_mag_aaral'],
    'age': ['age'],
    'sex': ['sex', 'gender'],
    'lrn': ['lrn', 'learner_reference_number', 'learner_reference_no'],
    'grade': ['grade', 'grade_level', 'gr'],
    'section': ['section', 'sec'],

    # Subject Grades
    # Filipino (Row 24)
    'filipino_t1': ['filipino_t1', 'filipino_q1', 'filipino1', 'fil_t1', 'fil_1'],
    'filipino_t2': ['filipino_t2', 'filipino_q2', 'filipino2', 'fil_t2', 'fil_2'],
    'filipino_t3': ['filipino_t3', 'filipino_q3', 'filipino3', 'fil_t3', 'fil_3'],
    'filipino_final': ['filipino_final', 'filipino_fg', 'fil_final'],
    'filipino_remarks': ['filipino_remarks', 'fil_remarks'],

    # English (Row 25)
    'english_t1': ['english_t1', 'english_q1', 'english1', 'eng_t1', 'eng_1'],
    'english_t2': ['english_t2', 'english_q2', 'english2', 'eng_t2', 'eng_2'],
    'english_t3': ['english_t3', 'english_q3', 'english3', 'eng_t3', 'eng_3'],
    'english_final': ['english_final', 'english_fg', 'eng_final'],
    'english_remarks': ['english_remarks', 'eng_remarks'],

    # Mathematics (Row 26)
    'math_t1': ['math_t1', 'math_q1', 'mathematics_t1', 'math1', 'math_1'],
    'math_t2': ['math_t2', 'math_q2', 'mathematics_t2', 'math2', 'math_2'],
    'math_t3': ['math_t3', 'math_q3', 'mathematics_t3', 'math3', 'math_3'],
    'math_final': ['math_final', 'mathematics_final', 'math_fg'],
    'math_remarks': ['math_remarks', 'mathematics_remarks'],

    # Science (Row 27)
    'science_t1': ['science_t1', 'science_q1', 'sci_t1', 'science1', 'sci_1'],
    'science_t2': ['science_t2', 'science_q2', 'sci_t2', 'science2', 'sci_2'],
    'science_t3': ['science_t3', 'science_q3', 'sci_t3', 'science3', 'sci_3'],
    'science_final': ['science_final', 'science_fg', 'sci_final'],
    'science_remarks': ['science_remarks', 'sci_remarks'],

    # AP (Row 28)
    'ap_t1': ['ap_t1', 'ap_q1', 'araling_panlipunan_t1', 'ap1', 'ap_1'],
    'ap_t2': ['ap_t2', 'ap_q2', 'araling_panlipunan_t2', 'ap2', 'ap_2'],
    'ap_t3': ['ap_t3', 'ap_q3', 'araling_panlipunan_t3', 'ap3', 'ap_3'],
    'ap_final': ['ap_final', 'araling_panlipunan_final', 'ap_fg'],
    'ap_remarks': ['ap_remarks', 'araling_panlipunan_remarks'],

    # Values Education (Row 29)
    'values_t1': ['values_t1', 'values_education_t1', 'esp_t1', 'esp_1', 'val_t1'],
    'values_t2': ['values_t2', 'values_education_t2', 'esp_t2', 'esp_2', 'val_t2'],
    'values_t3': ['values_t3', 'values_education_t3', 'esp_t3', 'esp_3', 'val_t3'],
    'values_final': ['values_final', 'values_fg', 'esp_final'],
    'values_remarks': ['values_remarks', 'esp_remarks'],

    # TLE / Creative Tech (Row 30)
    'tle_t1': ['tle_t1', 'tle_q1', 'tle1', 'tle_1', 'creative_tech_t1', 'tech_t1'],
    'tle_t2': ['tle_t2', 'tle_q2', 'tle2', 'tle_2', 'creative_tech_t2', 'tech_t2'],
    'tle_t3': ['tle_t3', 'tle_q3', 'tle3', 'tle_3', 'creative_tech_t3', 'tech_t3'],
    'tle_final': ['tle_final', 'tle_fg'],
    'tle_remarks': ['tle_remarks'],

    # MAPEH (Row 31)
    'mapeh_t1': ['mapeh_t1', 'mapeh_q1', 'mapeh1'],
    'mapeh_t2': ['mapeh_t2', 'mapeh_q2', 'mapeh2'],
    'mapeh_t3': ['mapeh_t3', 'mapeh_q3', 'mapeh3'],
    'mapeh_final': ['mapeh_final', 'mapeh_fg'],
    'mapeh_remarks': ['mapeh_remarks'],

    # Music & Arts (Row 32)
    'musicarts_t1': ['musicarts_t1', 'music_arts_t1', 'music_t1', 'arts_t1'],
    'musicarts_t2': ['musicarts_t2', 'music_arts_t2', 'music_t2', 'arts_t2'],
    'musicarts_t3': ['musicarts_t3', 'music_arts_t3', 'music_t3', 'arts_t3'],
    'musicarts_final': ['musicarts_final', 'music_arts_final'],
    'musicarts_remarks': ['musicarts_remarks', 'music_arts_remarks'],

    # PE & Health (Row 33)
    'pehealth_t1': ['pehealth_t1', 'pe_health_t1', 'pe_t1', 'health_t1', 'pe_and_health_t1'],
    'pehealth_t2': ['pehealth_t2', 'pe_health_t2', 'pe_t2', 'health_t2', 'pe_and_health_t2'],
    'pehealth_t3': ['pehealth_t3', 'pe_health_t3', 'pe_t3', 'health_t3', 'pe_and_health_t3'],
    'pehealth_final': ['pehealth_final', 'pe_health_final'],
    'pehealth_remarks': ['pehealth_remarks', 'pe_health_remarks'],

    # Research I (Row 34)
    'research1_t1': ['research1_t1', 'research_1_t1', 'research_t1', 'res1_t1', 'research_i_t1'],
    'research1_t2': ['research1_t2', 'research_1_t2', 'research_t2', 'res1_t2', 'research_i_t2'],
    'research1_t3': ['research1_t3', 'research_1_t3', 'research_t3', 'res1_t3', 'research_i_t3'],
    'research1_final': ['research1_final', 'research_1_final'],
    'research1_remarks': ['research1_remarks'],

    # Math of Investigation (Row 35)
    'mathinv_t1': ['mathinv_t1', 'math_inv_t1', 'moi_t1', 'math_of_investigation_t1'],
    'mathinv_t2': ['mathinv_t2', 'math_inv_t2', 'moi_t2', 'math_of_investigation_t2'],
    'mathinv_t3': ['mathinv_t3', 'math_inv_t3', 'moi_t3', 'math_of_investigation_t3'],
    'mathinv_final': ['mathinv_final', 'math_inv_final', 'moi_final'],
    'mathinv_remarks': ['mathinv_remarks'],

    # General Average (Row 37)
    'genavg_t1': ['genavg_t1', 'general_average_t1', 'gen_avg_t1'],
    'genavg_t2': ['genavg_t2', 'general_average_t2', 'gen_avg_t2'],
    'genavg_t3': ['genavg_t3', 'general_average_t3', 'gen_avg_t3'],
    'genavg_final': ['genavg_final', 'general_average_final', 'gen_avg_final'],
    'genavg_remarks': ['genavg_remarks', 'general_average_remarks'],

    # Attendance - Class Days (Left: C..M / Total N; Right: R..AB / Total AC)
    'days_jun': ['days_jun', 'class_days_jun', 'jun_days'],
    'days_jul': ['days_jul', 'class_days_jul', 'jul_days'],
    'days_aug': ['days_aug', 'class_days_aug', 'aug_days'],
    'days_sep': ['days_sep', 'class_days_sep', 'sep_days'],
    'days_oct': ['days_oct', 'class_days_oct', 'oct_days'],
    'days_nov': ['days_nov', 'class_days_nov', 'nov_days'],
    'days_dec': ['days_dec', 'class_days_dec', 'dec_days'],
    'days_jan': ['days_jan', 'class_days_jan', 'jan_days'],
    'days_feb': ['days_feb', 'class_days_feb', 'feb_days'],
    'days_mar': ['days_mar', 'class_days_mar', 'mar_days'],
    'days_apr': ['days_apr', 'class_days_apr', 'apr_days'],

    # Attendance - Present
    'present_jun': ['present_jun', 'days_present_jun', 'jun_present'],
    'present_jul': ['present_jul', 'days_present_jul', 'jul_present'],
    'present_aug': ['present_aug', 'days_present_aug', 'aug_present'],
    'present_sep': ['present_sep', 'days_present_sep', 'sep_present'],
    'present_oct': ['present_oct', 'days_present_oct', 'oct_present'],
    'present_nov': ['present_nov', 'days_present_nov', 'nov_present'],
    'present_dec': ['present_dec', 'days_present_dec', 'dec_present'],
    'present_jan': ['present_jan', 'days_present_jan', 'jan_present'],
    'present_feb': ['present_feb', 'days_present_feb', 'feb_present'],
    'present_mar': ['present_mar', 'days_present_mar', 'mar_present'],
    'present_apr': ['present_apr', 'days_present_apr', 'apr_present'],

    # Attendance - Absent
    'absent_jun': ['absent_jun', 'days_absent_jun', 'jun_absent'],
    'absent_jul': ['absent_jul', 'days_absent_jul', 'jul_absent'],
    'absent_aug': ['absent_aug', 'days_absent_aug', 'aug_absent'],
    'absent_sep': ['absent_sep', 'days_absent_sep', 'sep_absent'],
    'absent_oct': ['absent_oct', 'days_absent_oct', 'oct_absent'],
    'absent_nov': ['absent_nov', 'days_absent_nov', 'nov_absent'],
    'absent_dec': ['absent_dec', 'days_absent_dec', 'dec_absent'],
    'absent_jan': ['absent_jan', 'days_absent_jan', 'jan_absent'],
    'absent_feb': ['absent_feb', 'days_absent_feb', 'feb_absent'],
    'absent_mar': ['absent_mar', 'days_absent_mar', 'mar_absent'],
    'absent_apr': ['absent_apr', 'days_absent_apr', 'apr_absent'],

    # Comments & Transfer
    'remarks_t1': ['remarks_t1', 'teacher_comments_t1', 'comments_t1', 'term1_comments'],
    'remarks_t2': ['remarks_t2', 'teacher_comments_t2', 'comments_t2', 'term2_comments'],
    'remarks_t3': ['remarks_t3', 'teacher_comments_t3', 'comments_t3', 'term3_comments'],
    'admitted_to_grade': ['admitted_to_grade', 'admitted_grade', 'admit_grade'],
    'eligible_for': ['eligible_for', 'eligible_for_admission', 'eligible_admission_to_grade', 'eligible_grade'],
    'school_head': ['school_head', 'principal', 'approved_school_head'],
    'adviser_name': ['adviser_name', 'teacher_name', 'adviser', 'teacher_adviser_name'],
    'admitted_in': ['admitted_in', 'admit_in'],
    'transfer_date': ['transfer_date', 'date', 'cancellation_date'],
}

def normalize_key(s: Any) -> str:
    """Normalizes string for header matching (lowercase, underscores, no special chars)"""
    if s is None:
        return ""
    text = str(s).strip().lower()
    text = re.sub(r'[\'\"\,]', '', text)
    text = re.sub(r'[\s\-\.\/\(\)\:\_\[\]]+', '_', text)
    return text.strip('_')

def parse_input_workbook(file_bytes_or_path) -> List[Dict[str, Any]]:
    """
    Parses any uploaded Excel file into a list of standardized student dictionaries.
    """
    if isinstance(file_bytes_or_path, (bytes, bytearray)):
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes_or_path), data_only=True)
    else:
        wb = openpyxl.load_workbook(file_bytes_or_path, data_only=True)

    ws = wb.active
    
    header_row = 1
    for r in range(1, min(10, ws.max_row + 1)):
        non_empty = [ws.cell(r, c).value for c in range(1, ws.max_column + 1) if ws.cell(r, c).value is not None]
        if len(non_empty) >= 3:
            header_row = r
            break

    col_mapping = {}
    for c in range(1, ws.max_column + 1):
        val = ws.cell(header_row, c).value
        if val is None:
            continue
        raw_key = normalize_key(val)
        matched_field = None
        for std_field, aliases in HEADER_MAP.items():
            if raw_key == std_field or raw_key in aliases:
                matched_field = std_field
                break
        if matched_field:
            col_mapping[c] = matched_field
        else:
            col_mapping[c] = raw_key

    students = []
    for r in range(header_row + 1, ws.max_row + 1):
        row_data = {}
        has_content = False
        for c, field in col_mapping.items():
            val = ws.cell(r, c).value
            if val is not None:
                has_content = True
                row_data[field] = val
        
        if has_content and (row_data.get('name') or row_data.get('lrn')):
            students.append(row_data)

    return students

def sanitize_sheet_name(name: str, used_names: set) -> str:
    """Sanitize sheet name to max 31 valid Excel characters and ensure uniqueness"""
    clean = re.sub(r'[\:\\\/\?\*\[\]]', '', str(name)).strip()
    if not clean:
        clean = "Sheet"
    
    # If first 31 chars is unique, use it directly
    if clean[:31].lower() not in used_names:
        candidate = clean[:31]
        used_names.add(candidate.lower())
        return candidate

    # If duplicate, leave room for _1, _2 suffix (up to 31 chars)
    base = clean[:27]
    counter = 1
    while f"{base}_{counter}".lower() in used_names:
        counter += 1
    candidate = f"{base}_{counter}"
    used_names.add(candidate.lower())
    return candidate

def sanitize_last_name(name: Any) -> str:
    """Extract clean last name for sheet title and filename"""
    if not name:
        return "Record"
    clean = str(name).strip()
    if ',' in clean:
        clean = clean.split(',')[0].strip()
    else:
        parts = clean.split()
        if len(parts) > 1:
            clean = parts[-1].strip()
        elif parts:
            clean = parts[0].strip()
    clean = re.sub(r'[^A-Za-z0-9_\-]', '', clean)
    return clean or "Record"

def format_2up_tab_name(idx1: int, name1: str, idx2: Optional[int], name2: Optional[str], page_type: str, used_names: set) -> str:
    """
    Formats tab name e.g.
    '1-2 FRONT Navarro - Cedro' or '5 FRONT Perez'
    """
    last1 = sanitize_last_name(name1)
    if idx2 is not None and name2 is not None:
        last2 = sanitize_last_name(name2)
        raw_tab = f"{idx1}-{idx2} {page_type} {last1} - {last2}"
    else:
        raw_tab = f"{idx1} {page_type} {last1}"
    
    return sanitize_sheet_name(raw_tab, used_names)

def safe_num(val):
    """Safely converts to float/int or returns original value"""
    if val is None or val == "":
        return None
    try:
        f = float(val)
        return int(f) if f.is_integer() else round(f, 2)
    except (ValueError, TypeError):
        return val

def wrap_remarks_text(text: Any, width: int = 56, max_lines: int = 3) -> List[str]:
    """Wraps text at specified width (56 characters) across up to max_lines"""
    if not text:
        return []
    clean_str = str(text).strip()
    if not clean_str:
        return []
    lines = textwrap.wrap(clean_str, width=width, break_long_words=False)
    return lines[:max_lines]

def get_output_filenames(students: List[Dict[str, Any]]) -> Tuple[str, str, str]:
    """Generates clean filenames for FRONT, BACK, and complete ZIP bundle"""
    if not students:
        return "SF9_FRONT_Empty.xlsx", "SF9_BACK_Empty.xlsx", "SF9_Complete_Empty.zip"
    first_clean = sanitize_last_name(students[0].get('name', 'First'))
    last_clean = sanitize_last_name(students[-1].get('name', 'Last'))
    
    if len(students) == 1:
        front_name = f"SF9_FRONT_{first_clean}.xlsx"
        back_name = f"SF9_BACK_{first_clean}.xlsx"
        zip_name = f"SF9_Complete_{first_clean}.zip"
    else:
        front_name = f"SF9_FRONT_{first_clean}-{last_clean}.xlsx"
        back_name = f"SF9_BACK_{first_clean}-{last_clean}.xlsx"
        zip_name = f"SF9_Complete_{first_clean}-{last_clean}.zip"

    return front_name, back_name, zip_name

# Subject Rows in FRONT Template (Rows 24 to 35)
SUBJECT_ROWS_FRONT = {
    'filipino': 24,
    'english': 25,
    'math': 26,
    'science': 27,
    'ap': 28,
    'values': 29,
    'tle': 30,
    'mapeh': 31,
    'musicarts': 32,
    'pehealth': 33,
    'research1': 34,
    'mathinv': 35,
}

# Attendance Months in BACK Template
# Left card cols: C..M (Total N); Right card cols: R..AB (Total AC)
ATTENDANCE_MONTHS_BACK_LEFT = [
    ('jun', 'C'), ('jul', 'D'), ('aug', 'E'), ('sep', 'F'),
    ('oct', 'G'), ('nov', 'H'), ('dec', 'I'), ('jan', 'J'),
    ('feb', 'K'), ('mar', 'L'), ('apr', 'M')
]
ATTENDANCE_MONTHS_BACK_RIGHT = [
    ('jun', 'R'), ('jul', 'S'), ('aug', 'T'), ('sep', 'U'),
    ('oct', 'V'), ('nov', 'W'), ('dec', 'X'), ('jan', 'Y'),
    ('feb', 'Z'), ('mar', 'AA'), ('apr', 'AB')
]

def populate_front_card(ws, student: Dict[str, Any], side: str, global_overrides: Dict[str, Any]):
    """
    Populates one card (Left or Right) on the FRONT sheet.
    side: 'left' (cols B..H) or 'right' (cols K..Q)
    """
    is_left = (side == 'left')
    
    # 1. Profile coordinates
    sy_cell = 'D12' if is_left else 'M12'
    name_cell = 'C14' if is_left else 'L14'
    age_cell = 'F14' if is_left else 'O14'
    sex_cell = 'H14' if is_left else 'Q14'
    lrn_cell = 'C15' if is_left else 'L15'
    grade_cell = 'E15' if is_left else 'N15'
    section_cell = 'H15' if is_left else 'Q15'

    sy = global_overrides.get('school_year') or student.get('school_year')
    if sy:
        ws[sy_cell] = str(sy)
        ws[sy_cell].alignment = CENTER_ALIGN
    
    if student.get('name'):
        ws[name_cell] = str(student.get('name'))
        ws[name_cell].alignment = LEFT_CENTER_ALIGN

    if student.get('age') is not None:
        ws[age_cell] = safe_num(student.get('age'))
        ws[age_cell].alignment = CENTER_ALIGN

    if student.get('sex'):
        ws[sex_cell] = str(student.get('sex'))
        ws[sex_cell].alignment = CENTER_ALIGN

    if student.get('lrn'):
        ws[lrn_cell] = str(student.get('lrn'))
        ws[lrn_cell].alignment = CENTER_ALIGN

    if student.get('grade') is not None:
        ws[grade_cell] = safe_num(student.get('grade'))
        ws[grade_cell].alignment = CENTER_ALIGN

    if student.get('section'):
        ws[section_cell] = str(student.get('section'))
        ws[section_cell].alignment = CENTER_ALIGN

    # 2. Subject Grades (Rows 24 to 35)
    # Left: D(T1), E(T2), F(T3), G(Final), H(Remarks)
    # Right: M(T1), N(T2), O(T3), P(Final), Q(Remarks)
    t1_col = 'D' if is_left else 'M'
    t2_col = 'E' if is_left else 'N'
    t3_col = 'F' if is_left else 'O'
    fg_col = 'G' if is_left else 'P'
    rem_col = 'H' if is_left else 'Q'

    for subj, row in SUBJECT_ROWS_FRONT.items():
        t1 = safe_num(student.get(f'{subj}_t1'))
        t2 = safe_num(student.get(f'{subj}_t2'))
        t3 = safe_num(student.get(f'{subj}_t3'))
        fg = safe_num(student.get(f'{subj}_final'))
        rem = student.get(f'{subj}_remarks')

        if t1 is not None:
            ws[f'{t1_col}{row}'] = t1
            ws[f'{t1_col}{row}'].alignment = CENTER_ALIGN

        if t2 is not None:
            ws[f'{t2_col}{row}'] = t2
            ws[f'{t2_col}{row}'].alignment = CENTER_ALIGN

        if t3 is not None:
            ws[f'{t3_col}{row}'] = t3
            ws[f'{t3_col}{row}'].alignment = CENTER_ALIGN

        # Final Grade
        if fg is not None:
            ws[f'{fg_col}{row}'] = fg
        else:
            if subj == 'mapeh':
                # MAPEH row 31 averages Music (32) and PE (33)
                if t1 is None:
                    ws[f'{t1_col}{row}'] = f'=IF(COUNT({t1_col}32:{t1_col}33)>0,ROUND(AVERAGE({t1_col}32:{t1_col}33),0),"")'
                if t2 is None:
                    ws[f'{t2_col}{row}'] = f'=IF(COUNT({t2_col}32:{t2_col}33)>0,ROUND(AVERAGE({t2_col}32:{t2_col}33),0),"")'
                if t3 is None:
                    ws[f'{t3_col}{row}'] = f'=IF(COUNT({t3_col}32:{t3_col}33)>0,ROUND(AVERAGE({t3_col}32:{t3_col}33),0),"")'
                ws[f'{fg_col}{row}'] = f'=IF(COUNT({t1_col}{row}:{t3_col}{row})>0,ROUND(AVERAGE({t1_col}{row}:{t3_col}{row}),0),"")'
            else:
                ws[f'{fg_col}{row}'] = f'=IF(COUNT({t1_col}{row}:{t3_col}{row})>0,ROUND(AVERAGE({t1_col}{row}:{t3_col}{row}),0),"")'

        ws[f'{fg_col}{row}'].alignment = CENTER_ALIGN

        # Remarks
        if rem is not None:
            ws[f'{rem_col}{row}'] = str(rem)
        else:
            ws[f'{rem_col}{row}'] = f'=IF(ISNUMBER({fg_col}{row}),IF({fg_col}{row}>=75,"Passed","Failed"),"")'
        ws[f'{rem_col}{row}'].alignment = CENTER_ALIGN

    # 3. General Average (Row 37)
    core_t1 = f"{t1_col}24:{t1_col}31,{t1_col}34:{t1_col}35"
    core_t2 = f"{t2_col}24:{t2_col}31,{t2_col}34:{t2_col}35"
    core_t3 = f"{t3_col}24:{t3_col}31,{t3_col}34:{t3_col}35"
    core_fg = f"{fg_col}24:{fg_col}31,{fg_col}34:{fg_col}35"

    t1_avg = safe_num(student.get('genavg_t1'))
    t2_avg = safe_num(student.get('genavg_t2'))
    t3_avg = safe_num(student.get('genavg_t3'))
    fg_avg = safe_num(student.get('genavg_final'))
    rem_avg = student.get('genavg_remarks')

    ws[f'{t1_col}37'] = t1_avg if t1_avg is not None else f'=IF(COUNT({core_t1})>0,ROUND(AVERAGE({core_t1}),0),"")'
    ws[f'{t1_col}37'].alignment = CENTER_ALIGN

    ws[f'{t2_col}37'] = t2_avg if t2_avg is not None else f'=IF(COUNT({core_t2})>0,ROUND(AVERAGE({core_t2}),0),"")'
    ws[f'{t2_col}37'].alignment = CENTER_ALIGN

    ws[f'{t3_col}37'] = t3_avg if t3_avg is not None else f'=IF(COUNT({core_t3})>0,ROUND(AVERAGE({core_t3}),0),"")'
    ws[f'{t3_col}37'].alignment = CENTER_ALIGN

    ws[f'{fg_col}37'] = fg_avg if fg_avg is not None else f'=IF(COUNT({core_fg})>0,ROUND(AVERAGE({core_fg}),0),"")'
    ws[f'{fg_col}37'].alignment = CENTER_ALIGN

    ws[f'{rem_col}37'] = str(rem_avg) if rem_avg is not None else f'=IF(ISNUMBER({fg_col}37),IF({fg_col}37>=75,"Passed","Failed"),"")'
    ws[f'{rem_col}37'].alignment = CENTER_ALIGN

def populate_back_card(ws, student: Dict[str, Any], side: str, global_overrides: Dict[str, Any]):
    """
    Populates one card (Left or Right) on the BACK sheet.
    side: 'left' (cols B..N) or 'right' (cols Q..AC)
    """
    is_left = (side == 'left')

    # 0. Student Name Header in B1 (Left) or Q1 (Right)
    name_cell = 'B1' if is_left else 'Q1'
    student_name = student.get('name')
    if student_name:
        ws[name_cell] = str(student_name).strip()
        ws[name_cell].alignment = LEFT_CENTER_ALIGN
        ws[name_cell].font = NAME_HEADER_FONT

    # 1. Attendance Record (Rows 4, 5, 6)
    months_map = ATTENDANCE_MONTHS_BACK_LEFT if is_left else ATTENDANCE_MONTHS_BACK_RIGHT
    start_col = 'C' if is_left else 'R'
    end_col = 'M' if is_left else 'AB'
    total_col = 'N' if is_left else 'AC'

    # Class Days (Row 4)
    for mon, col_letter in months_map:
        val = student.get(f'days_{mon}')
        if val is not None:
            ws[f'{col_letter}4'] = safe_num(val)
            ws[f'{col_letter}4'].alignment = CENTER_ALIGN
    ws[f'{total_col}4'] = f"=SUM({start_col}4:{end_col}4)"
    ws[f'{total_col}4'].alignment = CENTER_ALIGN

    # Days Present (Row 5)
    for mon, col_letter in months_map:
        val = student.get(f'present_{mon}')
        if val is not None:
            ws[f'{col_letter}5'] = safe_num(val)
            ws[f'{col_letter}5'].alignment = CENTER_ALIGN
    ws[f'{total_col}5'] = f"=SUM({start_col}5:{end_col}5)"
    ws[f'{total_col}5'].alignment = CENTER_ALIGN

    # Days Absent (Row 6)
    for mon, col_letter in months_map:
        val = student.get(f'absent_{mon}')
        if val is not None:
            ws[f'{col_letter}6'] = safe_num(val)
            ws[f'{col_letter}6'].alignment = CENTER_ALIGN
    ws[f'{total_col}6'] = f"=SUM({start_col}6:{end_col}6)"
    ws[f'{total_col}6'].alignment = CENTER_ALIGN

    # 2. Teacher Remarks (Capped at 56 characters per line)
    # Term 1: Left rows 9, 10, 11 (C) | Right rows 9, 10, 11 (R)
    # Term 2: Left rows 12, 13 (C) | Right rows 12, 13 (R)
    # Term 3: Left rows 15, 16 (C) | Right rows 15, 16 (R)
    comm_col = 'C' if is_left else 'R'

    rem1_lines = wrap_remarks_text(student.get('remarks_t1'), width=56, max_lines=3)
    if len(rem1_lines) > 0: ws[f'{comm_col}9'] = rem1_lines[0]; ws[f'{comm_col}9'].alignment = LEFT_CENTER_ALIGN
    if len(rem1_lines) > 1: ws[f'{comm_col}10'] = rem1_lines[1]; ws[f'{comm_col}10'].alignment = LEFT_CENTER_ALIGN
    if len(rem1_lines) > 2: ws[f'{comm_col}11'] = rem1_lines[2]; ws[f'{comm_col}11'].alignment = LEFT_CENTER_ALIGN

    rem2_lines = wrap_remarks_text(student.get('remarks_t2'), width=56, max_lines=2)
    if len(rem2_lines) > 0: ws[f'{comm_col}12'] = rem2_lines[0]; ws[f'{comm_col}12'].alignment = LEFT_CENTER_ALIGN
    if len(rem2_lines) > 1: ws[f'{comm_col}13'] = rem2_lines[1]; ws[f'{comm_col}13'].alignment = LEFT_CENTER_ALIGN

    rem3_lines = wrap_remarks_text(student.get('remarks_t3'), width=56, max_lines=2)
    if len(rem3_lines) > 0: ws[f'{comm_col}15'] = rem3_lines[0]; ws[f'{comm_col}15'].alignment = LEFT_CENTER_ALIGN
    if len(rem3_lines) > 1: ws[f'{comm_col}16'] = rem3_lines[1]; ws[f'{comm_col}16'].alignment = LEFT_CENTER_ALIGN

    # 3. Certificate of Transfer & Eligibility
    # Left: Admitted to E31 (merged E31:F31), Eligible M31, Approved C33 (merged C33:F33), Adviser J33 (merged J33:M33), Admitted In C40 (merged C40:I40), Date L40 (merged L40:N40)
    # Right: Admitted to T31 (merged T31:U31), Eligible AB31, Approved R33 (merged R33:U33), Adviser Y33 (merged Y33:AB33), Admitted In R40 (merged R40:X40), Date AA40 (merged AA40:AC40)
    admit_cell = 'E31' if is_left else 'T31'
    eligible_cell = 'M31' if is_left else 'AB31'
    head_cell = 'C33' if is_left else 'R33'
    adviser_cell = 'J33' if is_left else 'Y33'
    admitted_in_cell = 'C40' if is_left else 'R40'
    date_cell = 'L40' if is_left else 'AA40'

    admit_grade = global_overrides.get('admitted_to_grade') or student.get('admitted_to_grade')
    if admit_grade is not None:
        ws[admit_cell] = safe_num(admit_grade)
        ws[admit_cell].alignment = CENTER_ALIGN

    eligible_grade = global_overrides.get('eligible_for') or student.get('eligible_for')
    if eligible_grade is not None:
        ws[eligible_cell] = safe_num(eligible_grade)
        ws[eligible_cell].alignment = CENTER_ALIGN

    school_head = global_overrides.get('school_head') or student.get('school_head')
    if school_head:
        ws[head_cell] = str(school_head)
        ws[head_cell].alignment = CENTER_ALIGN

    adviser = global_overrides.get('adviser_name') or student.get('adviser_name')
    if adviser:
        ws[adviser_cell] = str(adviser)
        ws[adviser_cell].alignment = CENTER_ALIGN

    admitted_in = global_overrides.get('admitted_in') or student.get('admitted_in')
    if admitted_in:
        ws[admitted_in_cell] = str(admitted_in)
        ws[admitted_in_cell].alignment = LEFT_CENTER_ALIGN

    transfer_date = global_overrides.get('transfer_date') or student.get('transfer_date')
    if transfer_date:
        ws[date_cell] = str(transfer_date)
        ws[date_cell].alignment = CENTER_ALIGN

def generate_sf9_front_workbook(
    students: List[Dict[str, Any]], 
    template_path: str = FRONT_TEMPLATE_PATH,
    global_overrides: Optional[Dict[str, Any]] = None
) -> Tuple[openpyxl.Workbook, str]:
    """
    Generates 2-Up FRONT SF9 workbook (2 students per sheet).
    Tabs named e.g. '1-2 FRONT Navarro - Cedro'
    """
    if not os.path.exists(template_path):
        raise FileNotFoundError(f"FRONT Template not found at {template_path}")

    wb = openpyxl.load_workbook(template_path)
    base_sheet = wb['SF9 FRONT'] if 'SF9 FRONT' in wb.sheetnames else wb.active
    used_names = set()
    global_overrides = global_overrides or {}

    front_name, _, _ = get_output_filenames(students)

    if not students:
        return wb, front_name

    # Step by 2
    for i in range(0, len(students), 2):
        s1 = students[i]
        s2 = students[i+1] if i+1 < len(students) else None

        ws = wb.copy_worksheet(base_sheet)
        tab_title = format_2up_tab_name(
            idx1=i+1,
            name1=s1.get('name', f"Student_{i+1}"),
            idx2=i+2 if s2 else None,
            name2=s2.get('name', f"Student_{i+2}") if s2 else None,
            page_type="FRONT",
            used_names=used_names
        )
        ws.title = tab_title

        # Populate Left Card
        populate_front_card(ws, s1, side='left', global_overrides=global_overrides)

        # Populate Right Card (if s2 exists)
        if s2:
            populate_front_card(ws, s2, side='right', global_overrides=global_overrides)

    wb.remove(base_sheet)
    return wb, front_name

def generate_sf9_back_workbook(
    students: List[Dict[str, Any]], 
    template_path: str = BACK_TEMPLATE_PATH,
    global_overrides: Optional[Dict[str, Any]] = None
) -> Tuple[openpyxl.Workbook, str]:
    """
    Generates 2-Up BACK SF9 workbook (2 students per sheet).
    Tabs named e.g. '1-2 BACK Navarro - Cedro'
    """
    if not os.path.exists(template_path):
        raise FileNotFoundError(f"BACK Template not found at {template_path}")

    wb = openpyxl.load_workbook(template_path)
    base_sheet = wb['SF9 BACK'] if 'SF9 BACK' in wb.sheetnames else wb.active
    used_names = set()
    global_overrides = global_overrides or {}

    _, back_name, _ = get_output_filenames(students)

    if not students:
        return wb, back_name

    # Step by 2
    for i in range(0, len(students), 2):
        s1 = students[i]
        s2 = students[i+1] if i+1 < len(students) else None

        ws = wb.copy_worksheet(base_sheet)
        ws.row_dimensions[1].hidden = False
        ws.row_dimensions[1].height = 20.0
        tab_title = format_2up_tab_name(
            idx1=i+1,
            name1=s1.get('name', f"Student_{i+1}"),
            idx2=i+2 if s2 else None,
            name2=s2.get('name', f"Student_{i+2}") if s2 else None,
            page_type="BACK",
            used_names=used_names
        )
        ws.title = tab_title

        # Populate Left Card
        populate_back_card(ws, s1, side='left', global_overrides=global_overrides)

        # Populate Right Card (if s2 exists)
        if s2:
            populate_back_card(ws, s2, side='right', global_overrides=global_overrides)

    wb.remove(base_sheet)
    return wb, back_name

def generate_sf9_zip_bundle(
    students: List[Dict[str, Any]], 
    global_overrides: Optional[Dict[str, Any]] = None
) -> Tuple[io.BytesIO, str]:
    """
    Generates both FRONT and BACK workbooks and bundles them into an in-memory ZIP archive.
    """
    front_wb, front_name = generate_sf9_front_workbook(students, global_overrides=global_overrides)
    back_wb, back_name = generate_sf9_back_workbook(students, global_overrides=global_overrides)
    _, _, zip_name = get_output_filenames(students)

    front_buf = io.BytesIO()
    front_wb.save(front_buf)
    front_buf.seek(0)

    back_buf = io.BytesIO()
    back_wb.save(back_buf)
    back_buf.seek(0)

    zip_buf = io.BytesIO()
    with zipfile.ZipFile(zip_buf, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr(front_name, front_buf.getvalue())
        zf.writestr(back_name, back_buf.getvalue())

    zip_buf.seek(0)
    return zip_buf, zip_name
